import asyncio
import base64
import json
import os
import signal
import shlex
import sys
from pathlib import Path

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.executors import claude_agent_sdk_runner
from app.runtime.sandbox import native_tool_app, native_tool_health_probe


@pytest.mark.parametrize("link_depth", range(5))
def test_native_result_mount_rejects_symlink_ancestors(tmp_path, link_depth):
    from app.platform.sandbox.claude_native_filesystem import prepare_claude_tool_results
    parts = [".claude-config", "projects", "ai-platform", "session-a", "tool-results"]
    parent = tmp_path.joinpath(*parts[:link_depth])
    parent.mkdir(parents=True, exist_ok=True)
    outside = tmp_path / "outside"
    outside.mkdir()
    (parent / parts[link_depth]).symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError, match="native_tool_results_path_invalid"):
        prepare_claude_tool_results(tmp_path, "session-a")


def test_native_filesystem_exposes_only_bound_cli_result_directory(tmp_path):
    from app.platform.sandbox.docker_native_filesystem import native_container_filesystem
    from app.sandbox.api import WORKSPACE_RUNTIME_PRIVATE_ROOTS
    layout = native_container_filesystem(str(tmp_path), "/workspace", tmp_path / "socket",
        read_only_paths={}, private_roots=WORKSPACE_RUNTIME_PRIVATE_ROOTS, sdk_session_id="session-a")
    native = ".claude-config/projects/ai-platform/session-a/tool-results"
    assert layout["volumes"][str(tmp_path / native)] == {"bind": "/workspace/" + native, "mode": "ro"}
    assert "mode=0711" in layout["tmpfs"]["/workspace/.claude-config"]
    assert "mode=000" in layout["tmpfs"]["/workspace/.ai-platform"]
    assert not any("other-session" in path for path in layout["volumes"])


@pytest.mark.asyncio
@pytest.mark.skipif(not os.environ.get("AI_PLATFORM_NATIVE_MCP_TEST_IMAGE"), reason="Requires a prepared native command image and local Docker")
async def test_real_cli_file_result_is_usable_in_docker_native_sidecar(monkeypatch, tmp_path):
    import docker
    from app.platform.sandbox.docker_native_filesystem import native_container_filesystem
    from app.sandbox.api import WORKSPACE_RUNTIME_PRIVATE_ROOTS
    from tests import test_claude_mcp_registration as native_check

    workspace = tmp_path / "workspace"
    workspace.mkdir(mode=0o777)
    workspace.chmod(0o777)
    socket_parent = tmp_path / "socket"
    socket_parent.mkdir()
    session_id = "5e6dfc4d-1b04-4d94-80a2-719f58746ae1"
    layout = native_container_filesystem(str(workspace), str(workspace), socket_parent,
        read_only_paths={}, private_roots=WORKSPACE_RUNTIME_PRIVATE_ROOTS, sdk_session_id=session_id)
    real_runner = claude_agent_sdk_runner.run_claude_agent_sdk
    async def run_bound(**kwargs):
        return await real_runner(**{**kwargs, "session_id": session_id})
    monkeypatch.setattr(claude_agent_sdk_runner, "run_claude_agent_sdk", run_bound)
    await native_check.test_platform_runner_uses_native_mcp_file_with_authorized_bash(
        monkeypatch, workspace, "http", True)
    relative = ".claude-config/projects/ai-platform/" + session_id
    saved = next((workspace / relative / "tool-results").glob("*.docx"))
    (workspace / ".claude-config/settings.json").write_text("synthetic-private-config")
    (workspace / relative / "transcript.jsonl").write_text("synthetic-private-transcript")
    other = workspace / ".claude-config/projects/ai-platform/other-session/tool-results"
    other.mkdir(parents=True)
    (other / "private.txt").write_text("other session")
    # The CLI runs on the host; the command runs under the sidecar's actual
    # Linux identity and mount layout. Broker/UDS behavior has separate tests.
    layout.update(user="10001:10001", cap_add=[])
    script = f"""
import os
from pathlib import Path
from docx import Document
p = Path({str(workspace)!r})
assert os.getuid() == os.getgid() == 10001
source = Path({str(saved)!r})
doc = Document(source)
assert [doc.paragraphs[0].text, doc.tables[0].cell(0,0).text] == ['native paragraph', 'native table']
for private in ['.claude-config/settings.json', {str(relative + '/transcript.jsonl')!r}, {str(other.relative_to(workspace) / 'private.txt')!r}]:
    assert not (p / private).exists()
try:
    source.write_bytes(b'tampered')
except OSError:
    pass
else:
    raise AssertionError('native result mount is writable')
(p / 'docker-parsed.txt').write_text('native paragraph|native table')
"""
    client = docker.from_env()
    try:
        output = client.containers.run(os.environ["AI_PLATFORM_NATIVE_MCP_TEST_IMAGE"],
            entrypoint=["python", "-c"], command=[script], network_mode="none", remove=True, **layout)
        assert output == b""
        assert (workspace / "docker-parsed.txt").read_text() == "native paragraph|native table"
    finally:
        client.close()


def test_native_tool_launcher_binds_then_drops_privileges_before_serving(monkeypatch):
    order, captured = [], {}
    class Listener:
        def __enter__(self): return self
        def __exit__(self, *args): order.append("closed")
        def fileno(self): return 17
    def bind(path):
        assert str(path) == native_tool_app.NATIVE_SOCKET_PATH
        order.append("bind")
        return Listener()
    def drop(uid, gid):
        assert (uid, gid) == (10001, 10001)
        order.append("drop")
    def serve(*args, **kwargs):
        order.append("serve")
        captured.update(kwargs)
    monkeypatch.delenv("AI_PLATFORM_NATIVE_TOOL_SOCKET", raising=False)
    monkeypatch.setattr(native_tool_app, "_bind_native_tool_socket", bind)
    monkeypatch.setattr(native_tool_app, "_drop_native_tool_privileges", drop)
    monkeypatch.setitem(sys.modules, "uvicorn", type("Server", (), {"run": staticmethod(serve)}))
    assert native_tool_app.main() == 0
    assert order == ["bind", "drop", "serve", "closed"]
    assert captured["fd"] == 17
    assert "uds" not in captured


def test_native_tool_launcher_never_serves_after_failed_privilege_drop(monkeypatch):
    from contextlib import nullcontext
    monkeypatch.setattr(native_tool_app, "_bind_native_tool_socket", lambda path: nullcontext())
    def reject(uid, gid):
        raise RuntimeError("native_tool_capabilities_not_dropped")
    monkeypatch.setattr(native_tool_app, "_drop_native_tool_privileges", reject)
    with pytest.raises(RuntimeError, match="native_tool_capabilities_not_dropped"):
        native_tool_app.main()


@pytest.mark.parametrize(
    ("euid", "egid", "groups", "configured_uid", "configured_gid", "accepted"),
    [
        (10001, 10001, [], 10001, 10001, True),
        (10001, 10001, [10001], 10001, 10001, True),
        (0, 10001, [], 10001, 10001, False),
        (10001, 0, [], 10001, 10001, False),
        (10001, 10001, [0], 10001, 10001, False),
        (10001, 10001, [10001, 20000], 10001, 10001, False),
        (0, 0, [], 0, 0, False),
    ],
)
def test_native_tool_identity_requires_fixed_uid_gid_and_bounded_supplementary_groups(
    monkeypatch,
    euid,
    egid,
    groups,
    configured_uid,
    configured_gid,
    accepted,
):
    monkeypatch.setattr(native_tool_app.os, "name", "posix")
    monkeypatch.setattr(native_tool_app.os, "geteuid", lambda: euid, raising=False)
    monkeypatch.setattr(native_tool_app.os, "getegid", lambda: egid, raising=False)
    monkeypatch.setattr(native_tool_app.os, "getgroups", lambda: groups, raising=False)

    if accepted:
        native_tool_app._require_native_tool_identity(configured_uid, configured_gid)
    else:
        with pytest.raises(RuntimeError, match="native_tool_identity_invalid"):
            native_tool_app._require_native_tool_identity(configured_uid, configured_gid)


def test_native_tool_non_dumpable_is_set_and_verified(monkeypatch):
    calls = []
    monkeypatch.setattr(native_tool_app.os, "name", "posix")
    monkeypatch.setattr(
        native_tool_app,
        "_linux_prctl",
        lambda option, argument=0: calls.append((option, argument)) or 0,
    )

    native_tool_app._set_process_non_dumpable()

    assert calls == [
        (native_tool_app._PR_SET_DUMPABLE, 0),
        (native_tool_app._PR_GET_DUMPABLE, 0),
    ]


@pytest.mark.parametrize("results", [[-1], [0, 1]])
def test_native_tool_non_dumpable_fails_closed_when_set_or_verification_fails(monkeypatch, results):
    outcomes = iter(results)
    monkeypatch.setattr(native_tool_app.os, "name", "posix")
    monkeypatch.setattr(native_tool_app, "_linux_prctl", lambda _option, _argument=0: next(outcomes))

    with pytest.raises(RuntimeError, match="native_tool_non_dumpable_failed"):
        native_tool_app._set_process_non_dumpable()


def test_native_tool_non_dumpable_fails_closed_when_prctl_is_unavailable(monkeypatch):
    monkeypatch.setattr(native_tool_app.os, "name", "posix")

    def unavailable(_option, _argument=0):
        raise RuntimeError("native_tool_non_dumpable_unavailable")

    monkeypatch.setattr(native_tool_app, "_linux_prctl", unavailable)

    with pytest.raises(RuntimeError, match="native_tool_non_dumpable_unavailable"):
        native_tool_app._set_process_non_dumpable()


@pytest.mark.asyncio
async def test_native_tool_lifespan_revalidates_identity(monkeypatch, tmp_path):
    workspace = tmp_path / "workspace"
    socket_parent = workspace / ".ai-platform"
    socket_parent.mkdir(parents=True)
    socket_path = socket_parent / "native-tool.sock"
    identities = []

    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_TOKEN", "x" * 32)
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_WORKSPACE", str(workspace))
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_SOCKET", str(socket_path))
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_UID", "10001")
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_GID", "10001")
    monkeypatch.setattr(
        native_tool_app,
        "_require_native_tool_identity",
        lambda uid, gid: identities.append((uid, gid)),
    )
    app = native_tool_app.create_native_tool_app()

    async with app.router.lifespan_context(app):
        pass

    assert identities == [(10001, 10001)]


def test_native_tool_health_requires_the_internal_token(monkeypatch, tmp_path):
    workspace = tmp_path / "workspace"
    socket_parent = workspace / ".ai-platform"
    socket_parent.mkdir(parents=True)
    socket_path = socket_parent / "native-tool.sock"
    token = "x" * 32

    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_TOKEN", token)
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_WORKSPACE", str(workspace))
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_SOCKET", str(socket_path))
    monkeypatch.setattr(native_tool_app, "_require_native_tool_identity", lambda _uid, _gid: None)

    with TestClient(native_tool_app.create_native_tool_app()) as client:
        assert client.get("/health").status_code == 403
        assert client.get(
            "/health",
            headers={native_tool_app.NATIVE_TOOL_AUTH_HEADER: "wrong-token"},
        ).status_code == 403
        response = client.get(
            "/health",
            headers={native_tool_app.NATIVE_TOOL_AUTH_HEADER: token},
        )

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_native_tool_health_probe_reads_secret_from_environment_without_output(monkeypatch, capsys):
    token = "probe-token-" + ("x" * 32)
    captured = {}

    class Response:
        status_code = 200

        @staticmethod
        def json():
            return {"status": "ok"}

    class Client:
        def __init__(self, **kwargs):
            captured["client"] = kwargs

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def get(self, path, *, headers):
            captured["request"] = (path, headers)
            return Response()

    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_TOKEN", token)
    monkeypatch.setenv(
        "AI_PLATFORM_NATIVE_TOOL_SOCKET",
        native_tool_health_probe.NATIVE_TOOL_CONTAINER_SOCKET,
    )
    monkeypatch.setattr(
        native_tool_health_probe.httpx,
        "HTTPTransport",
        lambda **kwargs: captured.setdefault("transport", kwargs),
    )
    monkeypatch.setattr(native_tool_health_probe.httpx, "Client", Client)

    assert native_tool_health_probe.main() == 0
    assert captured["transport"] == {
        "uds": native_tool_health_probe.NATIVE_TOOL_CONTAINER_SOCKET
    }
    assert captured["request"] == (
        "/health",
        {native_tool_app.NATIVE_TOOL_AUTH_HEADER: token},
    )
    assert capsys.readouterr() == ("", "")


def test_native_tool_health_probe_failure_never_emits_secret_or_path(monkeypatch, capsys):
    token = "probe-token-" + ("x" * 32)
    private_path = "/private/host/workspace/native-tool.sock"

    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_TOKEN", token)
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_SOCKET", private_path)

    assert native_tool_health_probe.main() == 1
    captured = capsys.readouterr()
    assert captured == ("", "")
    assert token not in captured.out + captured.err
    assert private_path not in captured.out + captured.err


def test_native_tool_listener_is_prebound_and_not_inherited_by_commands(monkeypatch, tmp_path):
    import stat
    import tempfile
    import os
    temporary = tempfile.TemporaryDirectory(prefix="native-", dir="/tmp")
    tmp_path = Path(temporary.name)
    path = tmp_path / "native-tool.sock"
    monkeypatch.setattr(native_tool_app, "NATIVE_SOCKET_PATH", str(path))
    monkeypatch.setattr(native_tool_app.os, "geteuid", lambda: 0)
    original_fstat = os.fstat
    def runtime_owner(fd):
        node = original_fstat(fd)
        return os.stat_result((*node[:4], 10001, 10001, *node[6:]))
    monkeypatch.setattr(native_tool_app.os, "fstat", runtime_owner)
    ownership = []
    monkeypatch.setattr(native_tool_app.os, "fchown", lambda fd, uid, gid: ownership.append((uid, gid)))
    monkeypatch.setattr(native_tool_app.os, "chown", lambda *args, **kwargs: None)
    with native_tool_app._bind_native_tool_socket(path) as listener:
        assert not listener.get_inheritable()
        assert stat.S_ISSOCK(path.stat().st_mode)
        assert stat.S_IMODE(path.stat().st_mode) == 0o660
        assert stat.S_IMODE(tmp_path.stat().st_mode) == 0o750
    assert ownership == [(0, 10001)]
    temporary.cleanup()


def test_native_tool_drop_clears_groups_ids_and_checks_capabilities(monkeypatch):
    calls = []
    monkeypatch.setattr(native_tool_app.os, "setgroups", lambda groups: calls.append(("groups", groups)))
    monkeypatch.setattr(native_tool_app.os, "setresgid", lambda *ids: calls.append(("gid", ids)), raising=False)
    monkeypatch.setattr(native_tool_app.os, "setresuid", lambda *ids: calls.append(("uid", ids)), raising=False)
    monkeypatch.setattr(native_tool_app, "_require_native_tool_identity", lambda *ids: None)
    monkeypatch.setattr(Path, "read_text", lambda self: "CapInh: 0\nCapPrm: 0\nCapEff: 0\nCapAmb: 0")
    monkeypatch.setattr(native_tool_app, "_set_process_non_dumpable", lambda: calls.append("non_dumpable"))
    native_tool_app._drop_native_tool_privileges(10001, 10001)
    assert calls == [("groups", []), ("gid", (10001, 10001, 10001)), ("uid", (10001, 10001, 10001)), "non_dumpable"]
    monkeypatch.setattr(Path, "read_text", lambda self: "CapInh: 0\nCapPrm: 1\nCapEff: 0\nCapAmb: 0")
    with pytest.raises(RuntimeError, match="native_tool_capabilities_not_dropped"):
        native_tool_app._drop_native_tool_privileges(10001, 10001)


@pytest.mark.asyncio
async def test_native_tool_command_uses_minimal_environment_and_process_isolation(monkeypatch, tmp_path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    captured = {}

    class Process:
        pid = 123
        returncode = 0
        stdout = None
        stderr = None

        async def wait(self):
            return 0

    async def create_subprocess_exec(*args, **kwargs):
        captured["args"] = args
        captured["kwargs"] = kwargs
        return Process()

    async def terminate(_process, *, include_orphans=False):
        captured["include_orphans"] = include_orphans

    async def terminate_uid(uid):
        captured["terminated_uid"] = uid

    monkeypatch.setattr(
        native_tool_app,
        "_require_native_tool_identity",
        lambda uid, gid: captured.update(identity=(uid, gid)),
    )
    monkeypatch.setattr(native_tool_app.asyncio, "create_subprocess_exec", create_subprocess_exec)
    monkeypatch.setattr(native_tool_app, "_terminate_process_group", terminate)
    monkeypatch.setattr(native_tool_app, "_terminate_uid_processes", terminate_uid)
    monkeypatch.setenv("AI_PLATFORM_SECRET", "must-not-reach-native-command")
    monkeypatch.setenv("ANTHROPIC_AUTH_TOKEN", "must-not-reach-native-command")

    result = await native_tool_app._run_command(
        command="printf safe",
        workspace=workspace,
        uid=10001,
        gid=10001,
        timeout_ms=5_000,
    )

    assert result.returncode == 0
    assert captured["args"] == ("/bin/bash", "--noprofile", "--norc", "-c", "printf safe")
    assert captured["kwargs"]["cwd"] == str(workspace)
    assert captured["kwargs"]["start_new_session"] is True
    assert captured["identity"] == (10001, 10001)
    assert all(key not in captured["kwargs"] for key in ("user", "group", "extra_groups"))
    assert captured["kwargs"]["umask"] == 0o077
    assert captured["include_orphans"] is True
    assert captured["terminated_uid"] == 10001
    assert captured["kwargs"]["env"] == {
        "HOME": str(workspace),
        "TMPDIR": str(workspace / ".native-skill-tmp"),
        "PATH": "/app/.venv/bin:/usr/local/bin:/usr/bin:/bin",
        "NODE_PATH": "/opt/ai-platform/document-dependencies/node_modules",
        "LANG": "C.UTF-8",
        "LC_ALL": "C.UTF-8",
        "PYTHONDONTWRITEBYTECODE": "1",
    }


@pytest.mark.asyncio
async def test_native_tool_pre_spawn_failure_diagnostic_is_fixed_and_sanitized(monkeypatch, tmp_path):
    workspace = tmp_path / "private-workspace"
    workspace.mkdir()
    secret_command = "printf private-command"
    secret_token = "private-token-value"

    async def fail_spawn(*_args, **_kwargs):
        raise OSError(f"spawn failed for {secret_command} at {workspace} using {secret_token}")

    monkeypatch.setattr(native_tool_app, "_require_native_tool_identity", lambda _uid, _gid: None)
    monkeypatch.setattr(native_tool_app.asyncio, "create_subprocess_exec", fail_spawn)
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_TOKEN", secret_token)

    with pytest.raises(RuntimeError) as exc_info:
        await native_tool_app._run_command(
            command=secret_command,
            workspace=workspace,
            uid=10001,
            gid=10001,
        )

    diagnostic = str(exc_info.value)
    assert diagnostic == "native_tool_stage=pre_spawn;class=failure"
    assert exc_info.value.__context__ is None
    assert exc_info.value.__cause__ is None
    assert all(value not in diagnostic for value in (secret_command, secret_token, str(workspace), "PATH"))


@pytest.mark.asyncio
async def test_native_tool_post_spawn_failure_diagnostic_is_fixed_and_sanitized(monkeypatch, tmp_path):
    workspace = tmp_path / "private-workspace"
    workspace.mkdir()
    secret_command = "printf private-command"
    secret_token = "private-token-value"
    cleanup_calls = []

    class Process:
        pid = 123
        returncode = None
        stdout = None
        stderr = None

        async def wait(self):
            raise OSError(f"wait failed for {secret_command} at {workspace} using {secret_token}")

    async def create_subprocess_exec(*_args, **_kwargs):
        return Process()

    async def terminate_group(process, *, include_orphans=False):
        cleanup_calls.append((process.pid, include_orphans))

    async def terminate_uid(uid):
        cleanup_calls.append((uid, "uid"))

    monkeypatch.setattr(native_tool_app, "_require_native_tool_identity", lambda _uid, _gid: None)
    monkeypatch.setattr(native_tool_app.asyncio, "create_subprocess_exec", create_subprocess_exec)
    monkeypatch.setattr(native_tool_app, "_terminate_process_group", terminate_group)
    monkeypatch.setattr(native_tool_app, "_terminate_uid_processes", terminate_uid)
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_TOKEN", secret_token)

    with pytest.raises(RuntimeError) as exc_info:
        await native_tool_app._run_command(
            command=secret_command,
            workspace=workspace,
            uid=10001,
            gid=10001,
        )

    diagnostic = str(exc_info.value)
    assert diagnostic == "native_tool_stage=post_spawn;class=failure"
    assert exc_info.value.__context__ is None
    assert exc_info.value.__cause__ is None
    assert all(value not in diagnostic for value in (secret_command, secret_token, str(workspace), "PATH"))
    assert cleanup_calls == [(123, False), (10001, "uid")]


@pytest.mark.asyncio
async def test_native_tool_cancellation_terminates_the_process_group(monkeypatch, tmp_path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    terminated = []

    class Process:
        pid = 123
        returncode = None
        stdout = None
        stderr = None

        async def wait(self):
            raise asyncio.CancelledError()

    async def create_subprocess_exec(*_args, **_kwargs):
        return Process()

    async def terminate(process, *, include_orphans=False):
        terminated.append((process.pid, include_orphans))

    async def terminate_uid(uid):
        terminated.append((uid, "uid"))

    monkeypatch.setattr(native_tool_app, "_require_native_tool_identity", lambda _uid, _gid: None)
    monkeypatch.setattr(native_tool_app.asyncio, "create_subprocess_exec", create_subprocess_exec)
    monkeypatch.setattr(native_tool_app, "_terminate_process_group", terminate)
    monkeypatch.setattr(native_tool_app, "_terminate_uid_processes", terminate_uid)

    with pytest.raises(asyncio.CancelledError):
        await native_tool_app._run_command(
            command="sleep 1",
            workspace=workspace,
            uid=10001,
            gid=10001,
        )

    assert terminated == [(123, False), (10001, "uid")]


def test_native_tool_process_sweep_finds_detached_active_uid_processes(tmp_path):
    proc_root = tmp_path / "proc"
    for process_id, uid, state in (
        ("101", 10001, "S"),
        ("102", 10001, "Z"),
        ("103", 10002, "S"),
        ("104", 10001, "S"),
    ):
        process_dir = proc_root / process_id
        process_dir.mkdir(parents=True)
        (process_dir / "status").write_text(
            f"Name:\tprobe\nState:\t{state} (probe)\nUid:\t{uid}\t{uid}\t{uid}\t{uid}\n",
            encoding="utf-8",
        )

    original_getpid = native_tool_app.os.getpid
    try:
        native_tool_app.os.getpid = lambda: 101
        assert native_tool_app._active_process_ids_for_uid(10001, proc_root=proc_root) == [104]
    finally:
        native_tool_app.os.getpid = original_getpid


@pytest.mark.asyncio
async def test_native_tool_process_sweep_escalates_to_kill(monkeypatch):
    observed = [[111, 222], [111, 222], [111]]
    killed = []

    monkeypatch.setattr(native_tool_app, "NATIVE_TOOL_TERMINATION_GRACE_SECONDS", 0)
    monkeypatch.setattr(native_tool_app.os, "getpid", lambda: 111)
    monkeypatch.setattr(
        native_tool_app,
        "_active_process_ids_for_uid",
        lambda _uid: observed.pop(0),
    )
    monkeypatch.setattr(native_tool_app.os, "kill", lambda pid, sig: killed.append((pid, sig)))

    await native_tool_app._terminate_uid_processes(10001)

    assert killed == [
        (222, signal.SIGTERM),
        (222, native_tool_app.NATIVE_TOOL_FORCE_KILL_SIGNAL),
    ]


@pytest.mark.asyncio
async def test_native_tool_disconnect_cancels_inflight_command(monkeypatch, tmp_path):
    cancelled = []

    async def run_command(**_kwargs):
        try:
            await asyncio.Event().wait()
        except asyncio.CancelledError:
            cancelled.append(True)
            raise

    class DisconnectedRequest:
        async def is_disconnected(self):
            return True

    monkeypatch.setattr(native_tool_app, "_run_command", run_command)

    with pytest.raises(HTTPException) as exc_info:
        await native_tool_app._execute_with_disconnect_cancellation(
            http_request=DisconnectedRequest(),
            command="sleep 60",
            timeout_ms=1_000,
            workspace=tmp_path,
            uid=10001,
            gid=10001,
        )

    assert exc_info.value.status_code == 499
    assert cancelled == [True]


def test_native_skill_workspace_paths_are_confined_and_proxy_carries_command_as_data(monkeypatch, tmp_path):
    workspace = tmp_path / "workspace"
    (workspace / "inputs").mkdir(parents=True)
    (workspace / "outputs" / "delivery").mkdir(parents=True)
    (workspace / ".ai-platform").mkdir()
    subject = {"workspace_contract": "ai-platform.skill-workspace.v1"}

    for incompatible_subject in (
        {},
        {"workspace_contract": "ai-platform.skill-workspace.v2"},
    ):
        for tool_name, tool_input in (
            ("Read", {"file_path": "inputs/source.docx"}),
            ("LS", {"path": "inputs"}),
            ("Glob", {"path": "inputs", "pattern": "**/*.xlsx"}),
            ("Grep", {"path": "inputs", "pattern": "TODO"}),
            ("Write", {"file_path": "outputs/delivery/report.pdf"}),
        ):
            assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
                incompatible_subject,
                tool_name,
                tool_input,
                workspace_root=workspace,
            )

    assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Read",
        {"file_path": "inputs/source.docx"},
        workspace_root=workspace,
    )
    assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Write",
        {"file_path": "outputs/delivery/report.pdf"},
        workspace_root=workspace,
    )
    assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Read",
        {"file_path": "../outside.txt"},
        workspace_root=workspace,
    )
    assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Read",
        {"file_path": str(workspace / ".ai-platform" / "native-tool.sock")},
        workspace_root=workspace,
    )
    assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Glob",
        {"path": ".", "pattern": "inputs/**/*.xlsx"},
        workspace_root=workspace,
    )
    assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Read",
        {"file_path": ".claude/skills/native-skill/SKILL.md"},
        workspace_root=workspace,
    )
    assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Glob",
        {"path": ".", "pattern": ".claude/skills/**/*.py"},
        workspace_root=workspace,
    )
    for allowed_input in (
        {"path": ".", "pattern": "**/*.xlsx"},
        {"path": "inputs", "pattern": "**/*.xlsx"},
        {"path": ".", "pattern": "reports/**/*.md"},
        {"path": ".", "pattern": ".CLAUDE/SKILLS/**/*.py"},
        {"path": ".", "pattern": "foo..bar"},
        {"path": ".", "pattern": "{inputs,outputs}/**/*.[pj]y"},
        {"path": ".", "pattern": "outputs/.*"},
        {"path": ".", "pattern": "reports/file?.md"},
    ):
        assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
            subject,
            "Glob",
            allowed_input,
            workspace_root=workspace,
        )
    for forbidden_input in (
        {"path": ".", "pattern": "/proc/**"},
        {"path": ".", "pattern": "C:/outside/**"},
        {"path": ".", "pattern": "{safe,/proc}/**"},
        {"path": ".", "pattern": "{safe,C:/outside}/**"},
        {"path": ".", "pattern": "@(safe|/proc)/**"},
        {"path": ".", "pattern": "@(safe|C:/outside)/**"},
        {"path": ".", "pattern": "@(safe|../.home)/**"},
        {"path": ".", "pattern": ".claude/skills/@(../settings|ok)/**"},
        {"path": ".", "pattern": ".claude/skills/.[.]/settings.json"},
        {"path": ".", "pattern": ".claude/skills/[.][.]/settings.json"},
        {"path": ".", "pattern": "safe/[.][.]/.home/**"},
        {"path": ".", "pattern": ".HOME/**"},
        {"path": ".", "pattern": ".claude/settings.json"},
        {"path": ".", "pattern": ".ai-platform/**"},
        {"path": ".", "pattern": ".home/**"},
        {"path": ".", "pattern": ".claude/skills/{../settings.json,ok.md}"},
        {"path": ".", "pattern": "inputs/{../.home,ok}/**"},
        {"path": "inputs", "pattern": "{../.home,ok}/**"},
        {"path": ".claude-config", "pattern": "**/*"},
    ):
        assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
            subject, "Glob", forbidden_input, workspace_root=workspace,
        ), forbidden_input
    for internal_root in (".claude-config", ".home", ".pins"):
        assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
            subject,
            "Read",
            {"file_path": str(workspace / internal_root / "private")},
            workspace_root=workspace,
        )
    assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Read",
        {"file_path": str(workspace / ".tmp" / "work" / "dump.txt")},
        workspace_root=workspace,
    )
    assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Read",
        {"file_path": ".claude/settings.json"},
        workspace_root=workspace,
    )
    assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
        subject,
        "Read",
        {"file_path": "outputs/.pins/private.txt"},
        workspace_root=workspace,
    )

    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_SOCKET", "/run/ai-platform-native/native-tool.sock")
    monkeypatch.setenv("AI_PLATFORM_NATIVE_TOOL_TOKEN", "x" * 32)
    command = "python scripts/run.py 'quoted; value'"
    proxied = claude_agent_sdk_runner._native_tool_proxy_input(
        {"command": command, "timeout": 5_000}
    )

    assert proxied is not None
    assert "native_tool_proxy.py" in proxied["command"]
    assert " -I " in proxied["command"]
    assert str(claude_agent_sdk_runner._NATIVE_TOOL_PROXY_SCRIPT) in proxied["command"]
    assert " -m " not in proxied["command"]
    assert command not in proxied["command"]
    proxy_args = shlex.split(proxied["command"])
    encoded = proxy_args[-2]
    assert base64.b64decode(encoded).decode("utf-8") == command
    assert proxy_args[-1] == "5000"
    assert proxied["timeout"] == 5_000
    assert claude_agent_sdk_runner._native_tool_proxy_input(
        {"command": command, "timeout": 600_001}
    ) is None


def test_workspace_search_results_filter_private_paths_and_bound_unlimited_grep(tmp_path):
    workspace = tmp_path / "workspace"
    (workspace / "inputs").mkdir(parents=True)
    (workspace / ".pins").mkdir()
    (workspace / "outputs" / ".ai-platform").mkdir(parents=True)
    (workspace / "inputs" / "public.py").write_text("public", encoding="utf-8")
    (workspace / ".pins" / "private.py").write_text("secret", encoding="utf-8")
    (workspace / "outputs" / ".ai-platform" / "state.json").write_text(
        "secret",
        encoding="utf-8",
    )

    filtered, changed = claude_agent_sdk_runner._filtered_workspace_search_output(
        {
            "mode": "files_with_matches",
            "filenames": [
                "inputs/public.py",
                ".pins/private.py",
                "outputs/.ai-platform/state.json",
            ],
            "numFiles": 3,
            "totalFiles": 3,
        },
        tool_name="Grep",
        tool_input={"pattern": "secret", "path": "."},
        workspace_root=workspace,
    )

    assert changed is True
    assert filtered == {
        "mode": "files_with_matches",
        "filenames": ["inputs/public.py", "outputs/.ai-platform/state.json"],
        "numFiles": 2,
    }

    filtered_count, changed = (
        claude_agent_sdk_runner._filtered_workspace_search_output(
            {
                "mode": "count",
                "filenames": [],
                "content": "inputs/public.py:2\n.pins/private.py:41",
                "numFiles": 2,
                "numMatches": 43,
                "totalFiles": 2,
                "totalLines": 43,
                "countIsComplete": True,
            },
            tool_name="Grep",
            tool_input={
                "pattern": "secret",
                "path": ".",
                "output_mode": "count",
            },
            workspace_root=workspace,
        )
    )
    assert changed is True
    assert filtered_count == {
        "mode": "count",
        "filenames": [],
        "content": (
            "inputs/public.py:2\n"
            "[ai-platform filtered or truncated search output]"
        ),
        "numFiles": 1,
        "numMatches": 2,
    }

    filtered_glob, changed = (
        claude_agent_sdk_runner._filtered_workspace_search_output(
            {
                "durationMs": 17,
                "numFiles": 2,
                "filenames": ["inputs/public.py", ".pins/private.py"],
                "truncated": False,
                "totalMatches": 2,
                "countIsComplete": True,
            },
            tool_name="Glob",
            tool_input={"pattern": "**/*.py", "path": "."},
            workspace_root=workspace,
        )
    )
    assert changed is True
    assert filtered_glob == {
        "durationMs": 17,
        "numFiles": 1,
        "filenames": ["inputs/public.py"],
        "truncated": True,
        "countIsComplete": False,
    }
    assert claude_agent_sdk_runner._bounded_workspace_search_input(
        "Grep",
        {"pattern": "TODO", "head_limit": 0},
    ) == {
        "pattern": "TODO",
        "head_limit": claude_agent_sdk_runner._WORKSPACE_GREP_MAX_HEAD_LIMIT,
    }
    assert claude_agent_sdk_runner._bounded_workspace_search_input(
        "Grep",
        {"pattern": "TODO", "head_limit": -1},
    ) is None

    filtered_text, changed = (
        claude_agent_sdk_runner._filtered_workspace_search_output(
            ".pins/private.py:1:secret",
            tool_name="Grep",
            tool_input={"pattern": "secret", "path": ".", "output_mode": "content"},
            workspace_root=workspace,
        )
    )
    assert changed is True
    assert "private.py" not in filtered_text

    public_content = "inputs/public.py:1:documentation mentioning .pins/private.py"
    kept_content, changed = claude_agent_sdk_runner._filtered_workspace_search_output(
        public_content, tool_name="Grep",
        tool_input={"pattern": "documentation", "path": ".", "output_mode": "content"},
        workspace_root=workspace,
    )
    assert kept_content == public_content
    assert changed is False

    oversized_list = ["inputs/public.py\n" * 2_000 for _ in range(3)]
    filtered_large, changed = (
        claude_agent_sdk_runner._filtered_workspace_search_output(
            oversized_list,
            tool_name="Glob",
            tool_input={"pattern": "**/*.py", "path": "."},
            workspace_root=workspace,
        )
    )
    assert changed is True
    assert (
        len(json.dumps(filtered_large, ensure_ascii=False).encode("utf-8"))
        <= claude_agent_sdk_runner._WORKSPACE_SEARCH_MAX_OUTPUT_BYTES
    )
    assert sum(
        item.count("inputs/public.py")
        for item in filtered_large
        if isinstance(item, str)
    ) == claude_agent_sdk_runner._WORKSPACE_SEARCH_MAX_OUTPUT_LINES


def test_workspace_search_projection_preserves_complete_sdk_output(tmp_path):
    workspace = tmp_path / "workspace"
    (workspace / "inputs").mkdir(parents=True)
    (workspace / "inputs" / "public.py").write_text("TODO", encoding="utf-8")
    response = {
        "durationMs": 4,
        "numFiles": 1,
        "filenames": ["inputs/public.py"],
        "truncated": False,
        "totalMatches": 1,
        "countIsComplete": True,
    }

    projected, changed = claude_agent_sdk_runner._filtered_workspace_search_output(
        response,
        tool_name="Glob",
        tool_input={"pattern": "**/*.py", "path": "."},
        workspace_root=workspace,
    )

    assert changed is False
    assert projected == response


def test_workspace_glob_budget_truncation_drops_sdk_totals(tmp_path):
    workspace = tmp_path / "workspace"
    (workspace / "inputs").mkdir(parents=True)
    filenames = [f"inputs/file-{index:04}.py" for index in range(1_001)]

    projected, changed = claude_agent_sdk_runner._filtered_workspace_search_output(
        {
            "durationMs": 5,
            "numFiles": len(filenames),
            "filenames": filenames,
            "truncated": False,
            "totalMatches": len(filenames),
            "countIsComplete": True,
        },
        tool_name="Glob",
        tool_input={"pattern": "**/*.py", "path": "."},
        workspace_root=workspace,
    )

    assert changed is True
    assert projected["filenames"] == filenames[:1_000]
    assert projected["numFiles"] == 1_000
    assert projected["truncated"] is True
    assert "totalMatches" not in projected
    assert projected["countIsComplete"] is False


def test_workspace_search_retains_sdk_truncation_and_limits_byte_output(tmp_path):
    workspace = tmp_path / "workspace"
    public_file = workspace / "inputs" / "public.py"
    public_file.parent.mkdir(parents=True)
    public_file.write_text("TODO", encoding="utf-8")

    sdk_truncated, changed = claude_agent_sdk_runner._filtered_workspace_search_output(
        {
            "durationMs": 1,
            "numFiles": 1,
            "filenames": ["inputs/public.py"],
            "truncated": True,
            "totalMatches": 1_200,
            "countIsComplete": True,
        },
        tool_name="Glob",
        tool_input={"pattern": "**/*.py", "path": "."},
        workspace_root=workspace,
    )
    assert changed is True
    assert sdk_truncated["truncated"] is True
    assert "totalMatches" not in sdk_truncated
    assert sdk_truncated["countIsComplete"] is False

    oversized_content = "\n".join(
        f"inputs/public.py:1:{'x' * 2_000}" for _ in range(200)
    )
    projected, changed = claude_agent_sdk_runner._filtered_workspace_search_output(
        {
            "mode": "content",
            "numFiles": 0,
            "filenames": [],
            "content": oversized_content,
            "numLines": 200,
            "totalLines": 200,
            "appliedLimit": 200,
            "appliedOffset": 3,
        },
        tool_name="Grep",
        tool_input={"pattern": "TODO", "path": ".", "output_mode": "content"},
        workspace_root=workspace,
    )

    assert changed is True
    assert isinstance(projected, dict)
    assert len(json.dumps(projected, ensure_ascii=False).encode("utf-8")) <= (
        claude_agent_sdk_runner._WORKSPACE_SEARCH_MAX_OUTPUT_BYTES
    )
    assert len(projected["content"].splitlines()) < 200
    assert "filtered or truncated" in projected["content"]
    assert projected["appliedLimit"] == 200
    assert projected["appliedOffset"] == 3
    assert "totalLines" not in projected


@pytest.mark.parametrize(
    ("mode", "response", "expected"),
    [
        (
            "files_with_matches",
            {
                "mode": "files_with_matches",
                "numFiles": 1,
                "filenames": ["inputs/public.py"],
                "appliedLimit": 5,
                "appliedOffset": 2,
            },
            {
                "mode": "files_with_matches",
                "numFiles": 1,
                "filenames": ["inputs/public.py"],
                "appliedLimit": 5,
                "appliedOffset": 2,
            },
        ),
        (
            "content",
            {
                "mode": "content",
                "numFiles": 0,
                "filenames": [],
                "content": "inputs/public.py:1:TODO",
                "numLines": 1,
                "totalLines": 9,
                "appliedLimit": 5,
                "appliedOffset": 2,
            },
            {
                "mode": "content",
                "numFiles": 0,
                "filenames": [],
                "content": "inputs/public.py:1:TODO",
                "numLines": 1,
                "appliedLimit": 5,
                "appliedOffset": 2,
            },
        ),
        (
            "count",
            {
                "mode": "count",
                "numFiles": 1,
                "filenames": [],
                "content": "inputs/public.py:2",
                "numMatches": 2,
                "appliedLimit": 5,
                "appliedOffset": 2,
            },
            {
                "mode": "count",
                "numFiles": 1,
                "filenames": [],
                "content": "inputs/public.py:2",
                "numMatches": 2,
                "appliedLimit": 5,
                "appliedOffset": 2,
            },
        ),
    ],
)
def test_grep_modes_keep_pagination_fields(mode, response, expected, tmp_path):
    workspace = tmp_path / "workspace"
    (workspace / "inputs").mkdir(parents=True)
    (workspace / "inputs" / "public.py").write_text("TODO", encoding="utf-8")

    projected, changed = claude_agent_sdk_runner._filtered_workspace_search_output(
        response,
        tool_name="Grep",
        tool_input={"pattern": "TODO", "path": ".", "output_mode": mode},
        workspace_root=workspace,
    )

    assert changed is True if mode == "content" else changed is False
    assert projected == expected


def test_workspace_search_hides_other_sessions_and_symlink_targets(tmp_path):
    workspace = tmp_path / "workspace"
    private = workspace / ".pins" / "private.py"
    private.parent.mkdir(parents=True)
    private.write_text("private", encoding="utf-8")
    public = workspace / "inputs" / "public.py"
    public.parent.mkdir(parents=True)
    public.write_text("public", encoding="utf-8")
    (public.parent / "private-link.py").symlink_to(private)

    session_root = workspace / ".claude-config" / "projects" / "ai-platform"
    current_result = session_root / "session-a" / "tool-results" / "result.md"
    current_result.parent.mkdir(parents=True)
    current_result.write_text("visible native result", encoding="utf-8")
    other_result = session_root / "session-b" / "tool-results" / "private.md"
    other_result.parent.mkdir(parents=True)
    other_result.write_text("other session", encoding="utf-8")

    projected, changed = claude_agent_sdk_runner._filtered_workspace_search_output(
        {
            "mode": "files_with_matches",
            "numFiles": 5,
            "filenames": [
                "inputs/public.py",
                "inputs/private-link.py",
                ".pins/private.py",
                ".claude-config/projects/ai-platform/session-a/tool-results/result.md",
                ".claude-config/projects/ai-platform/session-b/tool-results/private.md",
            ],
        },
        tool_name="Grep",
        tool_input={"pattern": "private", "path": "."},
        workspace_root=workspace,
        native_session_id="session-a",
    )

    assert changed is True
    assert projected["filenames"] == [
        "inputs/public.py",
        ".claude-config/projects/ai-platform/session-a/tool-results/result.md",
    ]
    assert projected["numFiles"] == 2


def test_staged_skill_paths_are_immutable_to_sdk_mutating_tools(tmp_path):
    workspace = tmp_path / "workspace"
    skill_file = workspace / ".claude" / "skills" / "native-skill" / "SKILL.md"
    skill_file.parent.mkdir(parents=True)
    skill_file.write_text("original", encoding="utf-8")
    (workspace / "inputs").mkdir()
    (workspace / "outputs" / "delivery").mkdir(parents=True)
    (workspace / "output").mkdir()
    subject = {"workspace_contract": "ai-platform.skill-workspace.v1"}

    mutation_inputs = {
        "Write": lambda path: {"file_path": path, "content": "tampered"},
        "Edit": lambda path: {"file_path": path, "old_string": "original", "new_string": "tampered"},
        "NotebookEdit": lambda path: {"notebook_path": path, "new_source": "tampered"},
    }
    staged_paths = (
        ".claude/skills/native-skill/SKILL.md",
        "./.claude/skills/native-skill/SKILL.md",
        ".claude\\skills\\native-skill\\SKILL.md",
        "outputs/delivery/../../.claude/skills/native-skill/SKILL.md",
        "outputs\\delivery\\..\\..\\.claude\\skills\\native-skill\\SKILL.md",
        str(skill_file),
        skill_file.as_posix(),
    )
    for tool_name, tool_input in mutation_inputs.items():
        for staged_path in staged_paths:
            assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
                subject,
                tool_name,
                tool_input(staged_path),
                workspace_root=workspace,
            )

        assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
            subject,
            tool_name,
            tool_input(f"outputs/delivery/{tool_name.lower()}.txt"),
            workspace_root=workspace,
        )
        assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
            subject,
            tool_name,
            tool_input(f"output/{tool_name.lower()}.txt"),
            workspace_root=workspace,
        )
        for workspace_path in (
            "draft.txt",
            "working/facts.json",
            "outputs/not-delivery.txt",
            "artifacts/report.docx",
            "tasks/facts.json",
            "review/draft.md",
        ):
            assert claude_agent_sdk_runner._workspace_path_parameters_authorized(
                subject,
                tool_name,
                tool_input(workspace_path),
                workspace_root=workspace,
            )
        for protected_path in (
            "CLAUDE.md",
            "CLAUDE.local.md",
            ".ai-platform-opensandbox-lease.json",
            "inputs/source.docx",
            ".ai-platform/state.json",
            ".claude/settings.json",
            ".claude-config/settings.json",
            ".home/.profile",
            ".pins/catalog.json",
            ".tmp/runtime.txt",
        ):
            assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
                subject,
                tool_name,
                tool_input(protected_path),
                workspace_root=workspace,
            )

    assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
        {},
        "Write",
        {"file_path": ".claude/skills/native-skill/SKILL.md", "content": "tampered"},
        workspace_root=workspace,
    )

    linked_skill = workspace / "outputs" / "delivery" / "staged-skill-link"
    linked_output = workspace / ".claude" / "skills" / "output-link"
    try:
        linked_skill.symlink_to(skill_file.parent, target_is_directory=True)
        linked_output.symlink_to(workspace / "outputs" / "delivery", target_is_directory=True)
    except OSError:
        return

    for tool_name, tool_input in mutation_inputs.items():
        assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
            subject,
            tool_name,
            tool_input(f"outputs/delivery/staged-skill-link/{tool_name.lower()}.txt"),
            workspace_root=workspace,
        )
        assert not claude_agent_sdk_runner._workspace_path_parameters_authorized(
            subject,
            tool_name,
            tool_input(f".claude/skills/output-link/{tool_name.lower()}.txt"),
            workspace_root=workspace,
        )
