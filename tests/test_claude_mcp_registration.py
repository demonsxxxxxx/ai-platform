import asyncio
import base64
import io
import json
import os
from pathlib import Path
import re
import shlex
import sys
from zipfile import ZIP_STORED, ZipFile
from contextlib import asynccontextmanager
from types import SimpleNamespace
from urllib.parse import urlsplit

import pytest
from mcp.client.session import ClientSession
from mcp.shared.memory import create_client_server_memory_streams
from mcp_types import BlobResourceContents, CallToolResult, EmbeddedResource, TextContent, Tool

from app.execution.infrastructure.claude_mcp import ClaudeMcpRegistration
from app.mcp.infrastructure import client as mcp_client
from app.mcp.infrastructure.catalog import _ValidatedDiscoveryTarget
from tests.support.mcp_protocol_peers import RAW_TOOL, SCHEMA, SDK_TOOL, local_mcp_peers


@asynccontextmanager
async def connected_client(server):
    async with create_client_server_memory_streams() as (client_streams, server_streams):
        client_read, client_write = client_streams
        server_read, server_write = server_streams
        server_task = asyncio.create_task(
            server.run(
                server_read,
                server_write,
                server.create_initialization_options(),
            )
        )
        try:
            async with ClientSession(client_read, client_write) as client:
                await client.initialize()
                yield client
        finally:
            if not server_task.done():
                server_task.cancel()
            await asyncio.gather(server_task, return_exceptions=True)


def subject(server="gateway", tool=RAW_TOOL, url="https://mcp.example/mcp"):
    return {
        "identity": f"mcp__{server}__{tool}", "mcp_server": server, "mcp_tool": tool,
        "registered": True, "declared": True, "active": True, "distributed": True,
        "identity_authorized": True, "object_authorized": True, "parameters_authorized": True,
        "risk_level": "low", "write_capable": False, "parameter_delegation": "external_mcp",
        "public_tool_label": "Sequence lookup",
        "mcp_server_config": {"type": "http", "url": url,
            "headers": {"JWT-Authorization": "Bearer synthetic.jwt", "X-Static-Key": "synthetic"}},
    }


def allow_local_peer(monkeypatch, url):
    async def target(endpoint):
        assert endpoint in {url + "/mcp", url + "/sse"}
        return _ValidatedDiscoveryTarget(endpoint, endpoint, urlsplit(url).netloc, "127.0.0.1")
    monkeypatch.setattr(mcp_client, "_validated_discovery_target", target)


@pytest.mark.parametrize("pairs", [
    [("gateway", "a.b"), ("gateway", "a_b")],
    [("gateway", "a:b"), ("gateway", "a_b")],
    [("server.name", "one"), ("server_name", "two")],
])
def test_sdk_name_collisions_get_stable_unique_aliases(pairs):
    subjects = [subject(server, tool) for server, tool in pairs]
    keyed = {item["identity"]: item for item in subjects}
    configs = {item["mcp_server"]: item["mcp_server_config"] for item in subjects}
    registration = ClaudeMcpRegistration(
        keyed, configs, session_factory=None, list_tools=None
    )
    reversed_registration = ClaudeMcpRegistration(
        dict(reversed(list(keyed.items()))), configs, session_factory=None, list_tools=None
    )

    aliases = list(registration.sdk_names.values())
    assert len(set(alias.casefold() for alias in aliases)) == len(pairs)
    assert all(re.fullmatch(r"mcp__[A-Za-z0-9_-]{1,57}", alias) for alias in aliases)
    assert registration.sdk_names == reversed_registration.sdk_names
    assert all(
        registration.canonical_identity(alias) == identity
        for identity, alias in registration.sdk_names.items()
    )
    for identity, (server, remote_tool) in registration._identity_components.items():
        assert registration.sdk_names[identity] == (
            f"mcp__{registration.server_aliases[server]}__"
            f"{registration._tool_aliases[(server, remote_tool)]}"
        )


def test_ambiguous_canonical_names_fail_before_registration():
    from app.executors.claude.capability_policy import _canonical_tool_policy_subjects
    with pytest.raises(ValueError, match="mcp_identity_collision"):
        _canonical_tool_policy_subjects([subject("a__b", "c"), subject("a", "b__c")])


@pytest.mark.asyncio
async def test_selected_bridge_preserves_unmarked_native_resource_result():
    raw = bytes(range(256)) * 8193
    item = subject()

    async def call_tool(_name, _arguments):
        return CallToolResult(content=[EmbeddedResource(resource=BlobResourceContents(
            uri="memory://synthetic-resource", mime_type="application/octet-stream",
            blob=__import__("base64").b64encode(raw).decode("ascii"),
        ))])

    registration = ClaudeMcpRegistration(
        {item["identity"]: item}, {}, session_factory=None, list_tools=None,
    )
    bridge = registration._server(
        "gateway", [(registration._tool_aliases[("gateway", RAW_TOOL)], Tool(name=RAW_TOOL, input_schema=SCHEMA), RAW_TOOL)],
        SimpleNamespace(call_tool=call_tool),
    )["instance"]
    async with connected_client(bridge) as client:
        result = await client.call_tool(registration._tool_aliases[("gateway", RAW_TOOL)], {"query": "synthetic"})
    assert result.is_error is False
    assert result.content[0].resource.blob == __import__("base64").b64encode(raw).decode("ascii")


@pytest.mark.asyncio
async def test_selected_bridge_uses_catalog_and_invocation_budgets(monkeypatch):
    configs = []
    item = subject()

    @asynccontextmanager
    async def session_factory(config):
        configs.append(config)
        yield SimpleNamespace(
            call_tool=lambda *_args: asyncio.sleep(0, result=CallToolResult(
                content=[TextContent(type="text", text="ok")]
            ))
        )

    async def list_tools(_session):
        return [Tool(name=RAW_TOOL, input_schema=SCHEMA)]

    registration = ClaudeMcpRegistration(
        {item["identity"]: item}, {"gateway": item["mcp_server_config"]},
        session_factory=session_factory, list_tools=list_tools,
    )
    options = SimpleNamespace()
    async with registration.activate(options):
        server = options.mcp_servers["gateway"]["instance"]
        async with connected_client(server) as client:
            await client.call_tool(registration._tool_aliases[("gateway", RAW_TOOL)], {})
    assert [config["max_response_bytes"] for config in configs] == [1024 * 1024, 16 * 1024 * 1024]
    assert [config["max_response_total_bytes"] for config in configs] == [1024 * 1024, None]

    def reject(reason):
        raise ValueError(reason)

    stream = mcp_client._BoundedMcpStream(
        None, event_stream=True, limit=configs[1]["max_response_bytes"],
        total_limit=configs[1]["max_response_total_bytes"], fail=reject,
    )
    frame = b"data: " + b"x" * (9 * 1024 * 1024) + b"\n\n"
    stream.check_chunk(frame)
    stream.check_chunk(frame)
    with pytest.raises(ValueError, match="response_too_large"):
        stream.check_chunk(b"x" * (16 * 1024 * 1024 + 1))


@pytest.mark.asyncio
@pytest.mark.parametrize("transport", ["http", "sse"])
async def test_real_mcp_transport_preserves_auth_names_results_and_cleanup(monkeypatch, transport):
    for name in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "http_proxy", "https_proxy", "all_proxy"):
        monkeypatch.setenv(name, "http://127.0.0.1:1")
    monkeypatch.setenv("NO_PROXY", "")
    monkeypatch.setenv("no_proxy", "")
    with local_mcp_peers() as peer:
        allow_local_peer(monkeypatch, peer.url)
        config = subject(url=peer.url + ("/mcp" if transport == "http" else "/sse"))["mcp_server_config"]
        config["type"] = transport
        async with asyncio.timeout(10):
            async with mcp_client.open_mcp_session(config, timeout_seconds=3) as session:
                tools = await mcp_client.list_mcp_tools(session)
                assert tools[0].input_schema == SCHEMA
                result = await session.call_tool(RAW_TOOL, {"query": "synthetic"})
                assert result.structured_content == {"sequence": 7}
                assert result.is_error is False
                peer.is_error = True
                failed = await session.call_tool(RAW_TOOL, {"query": "synthetic"})
                assert failed.is_error is True
        assert [call["name"] for call in peer.calls] == [RAW_TOOL, RAW_TOOL]
        assert all({key.lower(): value for key, value in headers.items()}.get("jwt-authorization") == "Bearer synthetic.jwt" for _, headers in peer.requests)
        assert all({key.lower(): value for key, value in headers.items()}.get("x-static-key") == "synthetic" for _, headers in peer.requests)
        if transport == "sse":
            assert await asyncio.to_thread(peer.sse_closed.wait, 2)


@pytest.mark.asyncio
@pytest.mark.parametrize("case", ["error", "success", "exception", "duplicate", "different_input"])
async def test_tool_error_result_is_bound_to_one_admitted_call(case):
    item = subject(server="files", tool="list_files")
    arguments = {"path": "restricted", "source": "public"}
    @asynccontextmanager
    async def session_factory(_config):
        async def call_tool(_name, _arguments):
            if case == "exception":
                raise TimeoutError("response was not received")
            return CallToolResult(content=[TextContent(type="text", text="access_denied")],
                                  is_error=case != "success")
        yield SimpleNamespace(call_tool=call_tool)
    async def list_tools(_session):
        return [Tool(name="list_files", input_schema={"type": "object"})]
    registration = ClaudeMcpRegistration(
        {item["identity"]: item}, {"files": item["mcp_server_config"]},
        session_factory=session_factory, list_tools=list_tools,
    )
    registration.track_tool_call(item["identity"], "call-a", arguments)
    if case == "duplicate":
        registration.track_tool_call(item["identity"], "call-b", arguments)
    options = SimpleNamespace()
    call_failure = None
    async with registration.activate(options):
        server = options.mcp_servers["files"]["instance"]
        async with connected_client(server) as client:
            if case == "exception":
                from mcp.shared.exceptions import MCPError

                with pytest.raises(MCPError) as raised:
                    await client.call_tool("list_files", arguments)
                call_failure = raised.value
            else:
                result = await client.call_tool(
                    "list_files",
                    arguments if case != "different_input" else {"path": "other"},
                )
    if case == "exception":
        assert call_failure.code == 0
        assert call_failure.message == "response was not received"
    else:
        assert result.is_error is (case != "success")
    assert registration.consume_tool_error_result(item["identity"], "call-a") is (case == "error")
    assert registration.consume_tool_error_result(item["identity"], "call-a") is False
    assert registration.consume_tool_error_result(item["identity"], "call-b") is False


@pytest.mark.asyncio
async def test_registration_filters_catalog_preserves_results_and_closes_on_cancel():
    item = subject()
    closed = []
    calls = []
    ready = asyncio.Event()

    @asynccontextmanager
    async def session_factory(_config):
        async def call_tool(name, arguments):
            calls.append((name, arguments))
            return CallToolResult(content=[TextContent(type="text", text="result")], structured_content={"n": 1}, is_error=True)
        try:
            yield SimpleNamespace(call_tool=call_tool)
        finally:
            closed.append(True)

    async def list_tools(_session):
        return [Tool(name=RAW_TOOL, input_schema=SCHEMA), Tool(name="unselected", input_schema={"type": "object"})]

    config = {"gateway": item["mcp_server_config"]}
    registration = ClaudeMcpRegistration({item["identity"]: item}, config, session_factory=session_factory, list_tools=list_tools)
    assert registration.sdk_names[item["identity"]] == SDK_TOOL
    assert registration.canonical_identity(SDK_TOOL) == item["identity"]

    async def consume():
        options = SimpleNamespace()
        async with registration.activate(options):
            server = options.mcp_servers[registration.server_aliases["gateway"]]["instance"]
            async with connected_client(server) as client:
                listed = await client.list_tools()
                assert [tool.name for tool in listed.tools] == [registration._tool_aliases[("gateway", RAW_TOOL)]]
                for name in (registration._tool_aliases[("gateway", RAW_TOOL)], "unselected"):
                    result = await client.call_tool(name, {"query": "synthetic"})
                    assert result.is_error is True
                    if name != "unselected":
                        assert result.structured_content == {"n": 1}
            ready.set()
            await asyncio.Event().wait()

    task = asyncio.create_task(consume())
    try:
        await asyncio.wait_for(ready.wait(), timeout=3)
    finally:
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    assert closed == [True, True]
    assert calls == [(RAW_TOOL, {"query": "synthetic"})]


@pytest.mark.asyncio
@pytest.mark.parametrize("discovery_phase", ["connect", "list"])
async def test_cancel_during_discovery_closes_owner_without_waiting_for_timeout(discovery_phase):
    entered = asyncio.Event()
    closed = asyncio.Event()
    owner = None

    @asynccontextmanager
    async def session_factory(_config):
        nonlocal owner
        owner = asyncio.current_task()
        try:
            if discovery_phase == "connect":
                entered.set()
                await asyncio.Event().wait()
            yield object()
        finally:
            assert asyncio.current_task() is owner
            closed.set()

    async def list_tools(_session):
        entered.set()
        await asyncio.Event().wait()

    item = subject()
    registration = ClaudeMcpRegistration(
        {item["identity"]: item}, {"gateway": item["mcp_server_config"]},
        session_factory=session_factory, list_tools=list_tools,
    )

    async def activate():
        async with registration.activate(SimpleNamespace()):
            pytest.fail("cancelled discovery must not activate")

    task = asyncio.create_task(activate())
    await asyncio.wait_for(entered.wait(), timeout=2)
    task.cancel()
    try:
        # No second cancellation: wait_for would mask slow cleanup by cancelling
        # the owner itself. The normal discovery timeout is ten seconds.
        done, _pending = await asyncio.wait({task}, timeout=1)
        assert done == {task}
        with pytest.raises(asyncio.CancelledError):
            await task
        assert closed.is_set()
    finally:
        if not task.done():
            task.cancel()
            await asyncio.gather(task, return_exceptions=True)


@pytest.mark.asyncio
@pytest.mark.parametrize("remote_error", [False, True])
@pytest.mark.parametrize("server_name", ["gateway", "server.name"])
@pytest.mark.parametrize("integration_kind", ["mcp", "enterprise_gateway"])
async def test_installed_claude_cli_selected_mcp_end_to_end(monkeypatch, tmp_path, remote_error, server_name, integration_kind, tool_search=False):
    from app.executors import claude_agent_sdk_runner as runner


    with local_mcp_peers() as peer:
        peer.is_error = remote_error
        peer.search_tools = tool_search
        selected_subject = {**subject(server=server_name, url=peer.url + "/mcp"), "write_capable": True}
        if integration_kind == "enterprise_gateway":
            selected_subject["mcp_server_config"].update(
                integration_kind=integration_kind, headers={"Authorization": "Bearer synthetic.jwt"},
            )
        selected_subject["mcp_server_config"]["credential_subject"] = {"tenant_id": "tenant", "user_id": "user"}
        credential_reads = []

        async def credential_resolver(server_id):
            credential_reads.append(server_id)
            return "synthetic.jwt"

        registration = ClaudeMcpRegistration(
            {selected_subject["identity"]: selected_subject}, {}, session_factory=None, list_tools=None
        )
        peer.sdk_tool = registration.sdk_names[selected_subject["identity"]]
        allow_local_peer(monkeypatch, peer.url)
        (tmp_path / ".mcp.json").write_text(json.dumps({
            "mcpServers": {"unconfigured": subject(server=server_name, url=peer.url + "/mcp")["mcp_server_config"]},
        }), encoding="utf-8")
        settings = SimpleNamespace(
            claude_agent_sdk_enabled=True, claude_agent_sdk_max_turns=4,
            claude_agent_sdk_timeout_seconds=40, claude_agent_sdk_skills="",
            claude_agent_permission_mode="dontAsk", claude_agent_model="claude-sonnet-4-6",
            anthropic_model="", anthropic_base_url=peer.url, anthropic_auth_token="synthetic-model-token",
            openai_api_key="",
        )
        monkeypatch.setattr(runner, "get_settings", lambda: settings)
        real_env = runner.build_sdk_env

        def synthetic_env(*, cwd, model_max_output_tokens=None):
            env = real_env(
                cwd=cwd,
                model_max_output_tokens=model_max_output_tokens,
            )
            env.update({"ANTHROPIC_BASE_URL": peer.url, "ANTHROPIC_AUTH_TOKEN": "synthetic-model-token",
                "CLAUDE_CODE_DISABLE_NONESSENTIAL_TRAFFIC": "1", "NO_PROXY": "127.0.0.1,localhost"})
            for name in ("TEMP", "TMP", "APPDATA", "LOCALAPPDATA"):
                if name in os.environ:
                    env[name] = os.environ[name]
            if tool_search == "default":
                env.pop("ENABLE_TOOL_SEARCH", None)
            elif tool_search:
                env["ENABLE_TOOL_SEARCH"] = "true"
            else:
                env["ENABLE_TOOL_SEARCH"] = "false"
            return env

        monkeypatch.setattr(runner, "build_sdk_env", synthetic_env)
        receipts = []
        events = []

        async def receipt(value):
            receipts.append(value)
            return True

        async def event(values):
            events.extend(values)
            return True

        result = await runner.run_claude_agent_sdk(
            prompt="Use the selected lookup once for synthetic, then report its result.",
            cwd=tmp_path, skill_id=None, skills=[], execution_policy="sandbox_brokered",
            tool_policy_subjects=[selected_subject],
            on_capability_evidence=receipt, on_agent_event=event,
            run_id="run_mcp_check", attempt_id="attempt_mcp_check",
            mcp_credential_resolver=credential_resolver,
        )
        assert credential_reads and set(credential_reads) == {server_name}
        assert result.error is None, result.runtime_diagnostics
        assert peer.models
        # Discovery uses catalog-only; invocation uses the bounded session.
        assert sum(method == "initialize" for method, _ in peer.requests) == 2
        if tool_search:
            assert any(tool["name"] == "ToolSearch" for tool in peer.models[0].get("tools", [])), peer.models[0].get("tools")
            assert any(tool.get("defer_loading") for tool in peer.models[0].get("tools", [])), peer.models[0].get("tools")
        else:
            assert all(
                {tool["name"] for tool in request.get("tools", [])}
                == {peer.sdk_tool, "mcp__ai-platform-response__attach_file"}
                for request in peer.models
            )
        tool_result_texts = [
            block.get("content", "")
            for request in peer.models
            for message in request.get("messages", [])
            for block in message.get("content", [])
            if isinstance(block, dict) and block.get("type") == "tool_result"
        ]
        if tool_search:
            search_results = [
                block for request in peer.models for message in request.get("messages", [])
                for block in message.get("content", [])
                if isinstance(block, dict) and block.get("tool_use_id") == "toolu_search"
            ]
            assert search_results and all(not block.get("is_error") for block in search_results), search_results
            assert peer.sdk_tool in str(search_results), search_results
            assert "unselected_sibling" not in str(peer.models)
        selected_tool = next(
            tool for request in peer.models for tool in request.get("tools", [])
            if tool["name"] == peer.sdk_tool
        )
        assert selected_tool["input_schema"] == SCHEMA
        if not remote_error:
            assert any(
                isinstance(content, str)
                and json.JSONDecoder().raw_decode(content)[0] == {"sequence": 7}
                for content in tool_result_texts
            ), tool_result_texts
        assert peer.calls == [{"name": RAW_TOOL, "arguments": {"query": "synthetic"}}], (
            result.turn_diagnostics, result.runtime_diagnostics,
            [block for request in peer.models for message in request.get("messages", [])
             for block in message.get("content", []) if isinstance(block, dict) and block.get("type") == "tool_result"],
        )
        if integration_kind == "enterprise_gateway":
            for method, request_headers in peer.requests:
                normalized = {key.lower(): value for key, value in request_headers.items()}
                assert normalized.get("authorization") == "Bearer synthetic.jwt"
                assert "jwt-authorization" not in normalized
        assert [(value["lifecycle_phase"], value["lifecycle_status"]) for value in receipts] == [
            ("invocation_requested", "invoking"), ("failed", "failed") if remote_error else ("completed", "succeeded"),
        ]
        assert all(value["canonical_identity"] == selected_subject["identity"] for value in receipts)
        assert receipts[-1]["trust_basis"] == (
            "tool_call_bound_error_result" if remote_error else "tool_call_bound_invocation"
        )

        assert {value.event_type for value in events} >= {"policy.allowed", "tool.started", "tool.failed" if remote_error else "tool.completed"}
        assert peer.sdk_tool not in str([value.payload for value in events])

        assert "synthetic.jwt" not in str([value.payload for value in events])


@pytest.mark.asyncio
@pytest.mark.parametrize("transport,authorized", [("http", True), ("sse", True), ("http", False)])
async def test_platform_runner_uses_native_mcp_file_with_authorized_bash(monkeypatch, tmp_path, transport, authorized):
    from docx import Document
    from app.executors import claude_agent_sdk_runner as runner
    from app.required_tool_contract import with_sandbox_local_tool_capability_subjects

    workspace = Path(str(tmp_path)[4:]) if str(tmp_path).startswith("\\\\?\\") else tmp_path
    document = Document()
    document.add_paragraph("native paragraph")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "native table"
    buffer = io.BytesIO()
    document.save(buffer)
    with ZipFile(buffer, "a", compression=ZIP_STORED) as archive:
        archive.writestr("opaque-padding.bin", b"P" * (1024 * 1024))
    raw = buffer.getvalue()
    assert len(raw) > 1024 * 1024
    parsed = tmp_path / "parsed.json"
    saved = {}
    receipts, lifecycle, events = [], [], []

    async def evidence(value):
        receipts.append(value)
        return True

    async def native_lifecycle(value):
        lifecycle.append(value)
        return True

    async def agent_events(values):
        events.extend(values)
        return True

    with local_mcp_peers() as peer:
        peer.tools = [{"name": "get_file", "description": "Return an ordinary file resource",
                       "inputSchema": {"type": "object", "properties": {}}}]
        peer.tool_result = {"content": [{"type": "resource", "resource": {
            "uri": "memory://sample.docx",
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "blob": base64.b64encode(raw).decode("ascii"),
        }}], "isError": False}
        item = subject(server="files", tool="get_file", url=peer.url + ("/sse" if transport == "sse" else "/mcp"))
        item["identity_authorized"] = authorized
        item["mcp_server_config"]["type"] = transport
        registration = ClaudeMcpRegistration({item["identity"]: item}, {}, session_factory=None, list_tools=None)
        peer.sdk_tool = registration.sdk_names[item["identity"]]
        allow_local_peer(monkeypatch, peer.url)

        def model_reply(request):
            results = {
                block.get("tool_use_id"): block
                for message in request.get("messages", [])
                for block in message.get("content", [])
                if isinstance(block, dict) and block.get("type") == "tool_result"
            }
            if "toolu_file" not in results:
                return {"type": "tool_use", "id": "toolu_file", "name": peer.sdk_tool, "input": {}}
            if not authorized or "toolu_parse" in results:
                return {"type": "text", "text": "Native file check complete."}
            content = results["toolu_file"].get("content", [])
            text = content if isinstance(content, str) else "\n".join(
                block.get("text", "") for block in content if isinstance(block, dict)
            )
            match = re.search(r"saved to (.+?\.docx)(?:\n|$)", text)
            assert match, text
            source = os.path.abspath(match.group(1))
            config_root = os.path.abspath(workspace / ".claude-config")
            assert os.path.commonpath([source, config_root]) == config_root
            if os.name == "nt" and not source.startswith("\\\\?\\"):
                source = "\\\\?\\" + source
            saved["path"] = source
            source_argument = Path(source).as_posix()
            script = (
                "import io,json; from pathlib import Path; from docx import Document; "
                f"d=Document(io.BytesIO(Path({source_argument!r}).read_bytes())); "
                f"Path({parsed.as_posix()!r}).write_text(json.dumps([d.paragraphs[0].text,d.tables[0].cell(0,0).text]),encoding='utf-8')"
            )
            executable = sys.executable.replace("\\", "/")
            command = f"{shlex.quote(executable)} -c {shlex.quote(script)}"
            return {"type": "tool_use", "id": "toolu_parse", "name": "Bash", "input": {"command": command}}

        peer.model_reply = model_reply
        settings = SimpleNamespace(
            claude_agent_sdk_enabled=True, claude_agent_sdk_max_turns=5,
            claude_agent_sdk_timeout_seconds=45, claude_agent_sdk_skills="",
            claude_agent_permission_mode="dontAsk", claude_agent_model="claude-sonnet-4-6",
            anthropic_model="", anthropic_base_url=peer.url,
            anthropic_auth_token="synthetic-model-token", openai_api_key="",
        )
        monkeypatch.setattr(runner, "get_settings", lambda: settings)
        real_env = runner.build_sdk_env

        def synthetic_env(*, cwd, model_max_output_tokens=None):
            env = real_env(cwd=cwd, model_max_output_tokens=model_max_output_tokens)
            env.update({"CLAUDE_CODE_DISABLE_NONESSENTIAL_TRAFFIC": "1",
                        "NO_PROXY": "127.0.0.1,localhost", "ENABLE_TOOL_SEARCH": "false"})
            for key in ("TEMP", "TMP", "APPDATA", "LOCALAPPDATA"):
                if key in os.environ:
                    env[key] = os.environ[key]
            return env

        monkeypatch.setattr(runner, "build_sdk_env", synthetic_env)
        subjects = with_sandbox_local_tool_capability_subjects(
            [item], sandbox_provider="opensandbox",
            authorized_sandbox_tool_identities={"Bash"} if authorized else set(),
        )
        result = await runner.run_claude_agent_sdk(
            prompt="Get the selected file, then use authorized Bash to read its paragraph and table.",
            cwd=workspace, skill_id=None, skills=[], execution_policy="sandbox_brokered",
            tool_policy_subjects=subjects, on_capability_evidence=evidence,
            on_agent_event=agent_events, on_tool_lifecycle=native_lifecycle,
            run_id="run_native_file", attempt_id="attempt_native_file",
        )
        if not authorized:
            assert not peer.calls
            assert not parsed.exists()
            assert not saved
            return
        assert result.error is None, result.runtime_diagnostics
        assert peer.calls == [{"name": "get_file", "arguments": {}}]
        assert Path(saved["path"]).read_bytes() == raw
        assert json.loads(parsed.read_text(encoding="utf-8")) == ["native paragraph", "native table"]
        assert lifecycle
        selected_receipts = [value for value in receipts if value["canonical_identity"] == item["identity"]]
        assert [value["lifecycle_phase"] for value in selected_receipts] == ["invocation_requested", "completed"]
        assert "synthetic.jwt" not in str([value.payload for value in events])
        assert ".claude-config" not in str([value.payload for value in events])


@pytest.mark.asyncio
@pytest.mark.parametrize("transport", ["http", "sse"])
@pytest.mark.parametrize("integration_kind", ["mcp", "enterprise_gateway"])
async def test_bound_executor_registration_refreshes_credentials_without_platform_store(monkeypatch, transport, integration_kind):
    from app.bootstrap.claude_mcp import prepare_claude_mcp
    from app.mcp.infrastructure import runtime

    def no_platform_store():
        raise AssertionError("executor must not read platform Redis")

    monkeypatch.setattr(runtime, "get_mcp_principal_jwt_store", no_platform_store)
    refreshed = []

    async def resolve(server_id):
        assert server_id == "gateway"
        jwt = f"synthetic-refreshed-{len(refreshed)}"
        refreshed.append(jwt)
        return jwt

    with local_mcp_peers() as peer:
        allow_local_peer(monkeypatch, peer.url)
        item = subject(url=peer.url + ("/sse" if transport == "sse" else "/mcp"))
        config = item["mcp_server_config"]
        header = "Authorization" if integration_kind == "enterprise_gateway" else "JWT-Authorization"
        config.update(type=transport, integration_kind=integration_kind,
                      headers={header: "Bearer synthetic-initial"},
                      credential_subject={"tenant_id": "tenant", "user_id": "user"})
        registration = prepare_claude_mcp(
            {item["identity"]: item}, {"gateway": config}, credential_resolver=resolve,
        )
        options = SimpleNamespace()
        async with registration.activate(options):
            assert registration.unavailable == {}
            server = options.mcp_servers[registration.server_aliases["gateway"]]["instance"]
            async with connected_client(server) as client:
                listed = await client.list_tools()
                assert len(listed.tools) == 1
                result = await client.call_tool(
                    listed.tools[0].name,
                    {"query": "synthetic"},
                )
                assert result.is_error is False
        values = [dict((key.lower(), value) for key, value in headers.items())[header.lower()]
                  for _method, headers in peer.requests]
        assert len(refreshed) >= len(values) >= 4
        assert len(set(values)) == len(values)
        assert all(value.removeprefix("Bearer ") in refreshed for value in values)
        assert peer.calls == [{"name": RAW_TOOL, "arguments": {"query": "synthetic"}}]


@pytest.mark.asyncio
async def test_registration_keeps_optional_server_absent_when_no_selected_tool_is_available():
    item = subject()
    closed = []
    activated = []

    @asynccontextmanager
    async def session_factory(_config):
        try:
            yield SimpleNamespace()
        finally:
            closed.append(True)

    async def list_tools(_session):
        return []

    registration = ClaudeMcpRegistration(
        {item["identity"]: item}, {"gateway": item["mcp_server_config"]},
        session_factory=session_factory, list_tools=list_tools,
    )
    options = SimpleNamespace()
    async with registration.activate(options):
        activated.append(True)
        assert "gateway" not in options.mcp_servers
    assert activated == [True]
    assert closed == [True]
    assert registration.unavailable == {"gateway": "selected_tool_unavailable"}


@pytest.mark.asyncio
async def test_registration_keeps_available_selected_tools_when_one_is_missing():
    available = subject(tool="lookup")
    missing = subject(tool="missing")
    closed = []

    @asynccontextmanager
    async def session_factory(_config):
        try:
            yield SimpleNamespace(
                call_tool=lambda name, arguments: asyncio.sleep(0, result=CallToolResult(
                    content=[TextContent(type="text", text=f"{name}:{arguments['query']}")]
                ))
            )
        finally:
            closed.append(True)

    async def list_tools(_session):
        return [Tool(name="lookup", input_schema={"type": "object"})]

    registration = ClaudeMcpRegistration(
        {available["identity"]: available, missing["identity"]: missing},
        {"gateway": available["mcp_server_config"]},
        session_factory=session_factory,
        list_tools=list_tools,
    )
    options = SimpleNamespace()
    async with registration.activate(options):
        assert registration.unavailable == {"gateway": "selected_tool_unavailable"}
        assert registration.server_aliases["gateway"] in options.mcp_servers
        server = options.mcp_servers[registration.server_aliases["gateway"]]["instance"]
        async with connected_client(server) as client:
            listed = await client.list_tools()
            assert [tool.name for tool in listed.tools] == [
                registration._tool_aliases[("gateway", "lookup")]
            ]
    assert closed == [True, True]


@pytest.mark.asyncio
async def test_registration_keeps_healthy_server_when_another_is_offline_and_cleans_up_in_owner():
    healthy = subject(server="healthy", tool="lookup")
    offline = subject(server="offline", tool="lookup")
    configs = {"healthy": {"server": "healthy"}, "offline": {"server": "offline"}}
    events = []
    call_names = []
    task_owners = {}
    discovered = set()
    both_discovered = asyncio.Event()

    @asynccontextmanager
    async def session_factory(config):
        server = config["server"]
        task = asyncio.current_task()
        task_owners[server] = [task, task]
        events.append((server, "open"))
        try:
            async def call_tool(name, arguments):
                call_names.append((server, name, arguments))
                return CallToolResult(content=[TextContent(type="text", text="ok")])

            yield SimpleNamespace(server=server, call_tool=call_tool)
        finally:
            task_owners[server][1] = asyncio.current_task()
            events.append((server, "close"))

    async def list_tools(session):
        discovered.add(session.server)
        if discovered == {"healthy", "offline"}:
            both_discovered.set()
        await asyncio.wait_for(both_discovered.wait(), timeout=0.5)
        if session.server == "offline":
            raise RuntimeError("offline")
        return [Tool(name="lookup", input_schema={"type": "object"})]

    registration = ClaudeMcpRegistration(
        {healthy["identity"]: healthy, offline["identity"]: offline},
        configs,
        session_factory=session_factory,
        list_tools=list_tools,
    )
    options = SimpleNamespace()
    async with registration.activate(options):
        assert registration.unavailable == {"offline": "connection_failed"}
        assert "offline" not in options.mcp_servers
        assert options.mcp_servers[registration.server_aliases["healthy"]]["type"] == "sdk"
        assert options.mcp_servers[registration.server_aliases["healthy"]]["type"] == "sdk"
        assert ("offline", "close") in events
        assert events.count(("healthy", "close")) == 1

        server = options.mcp_servers[registration.server_aliases["healthy"]]["instance"]
        async with connected_client(server) as client:
            listed = await client.list_tools()
            alias = registration._tool_aliases[("healthy", "lookup")]
            assert [tool.name for tool in listed.tools] == [alias]
            result = await client.call_tool(alias, {"query": "synthetic"})
            assert result.is_error is False

    assert events.index(("offline", "close")) < events.index(("healthy", "close"))
    assert task_owners["healthy"][0] is task_owners["healthy"][1]
    assert task_owners["offline"][0] is task_owners["offline"][1]
    assert call_names == [("healthy", "lookup", {"query": "synthetic"})]


@pytest.mark.asyncio
async def test_registration_timeout_does_not_limit_healthy_session_lifetime(monkeypatch):
    item = subject(tool="lookup")
    closed = []

    @asynccontextmanager
    async def session_factory(_config):
        try:
            yield SimpleNamespace(call_tool=lambda *_args: None)
        finally:
            closed.append(True)

    async def list_tools(_session):
        return [Tool(name="lookup", input_schema={"type": "object"})]

    from app.execution.infrastructure import claude_mcp

    monkeypatch.setattr(claude_mcp, "_MCP_DISCOVERY_TIMEOUT_SECONDS", 0.05)
    registration = ClaudeMcpRegistration(
        {item["identity"]: item}, {"gateway": item["mcp_server_config"]},
        session_factory=session_factory, list_tools=list_tools,
    )
    options = SimpleNamespace()
    async with registration.activate(options):
        await asyncio.sleep(0.1)
        assert registration.server_aliases["gateway"] in options.mcp_servers
        assert closed == [True]
    assert closed == [True, True]


@pytest.mark.asyncio
@pytest.mark.parametrize("code", [-32602, 408, -(2**53 - 1), 2**53 - 1, True, "408", {}, 2**53])
async def test_typed_session_failure_is_content_free_and_reraised_unchanged(code):
    from mcp.shared.exceptions import MCPError
    from mcp_types import CallToolRequestParams, ErrorData

    failure = MCPError(code=0, message="PRIVATE_MESSAGE", data={"secret": "PRIVATE_DATA"})
    failure.error = ErrorData.model_construct(
        code=code, message="PRIVATE_MESSAGE", data={"secret": "PRIVATE_DATA"}
    )
    observations = []
    callbacks = []

    def track(callback):
        callbacks.append(callback)
        return callback

    async def call_tool(name, arguments):
        raise failure

    registration = ClaudeMcpRegistration(
        {}, {}, session_factory=None, list_tools=None,
        callback_wrapper=track, failure_observer=observations.append,
    )
    registration._server(
        "gateway",
        [("lookup", Tool(name="lookup", input_schema={}), "lookup")],
        SimpleNamespace(call_tool=call_tool),
    )
    with pytest.raises(MCPError) as raised:
        await callbacks[0](
            SimpleNamespace(),
            CallToolRequestParams(name="lookup", arguments={"query": "PRIVATE_ARGUMENT"}),
        )
    assert raised.value is failure
    expected = {"source": "mcp_session_call", "scope": "run_attempt", "exception_type": "MCPError"}
    if type(code) is int and abs(code) <= 2**53 - 1:
        expected["code"] = code
    assert observations == [expected]
    assert "PRIVATE" not in json.dumps(observations)


@pytest.mark.asyncio
async def test_lowlevel_server_returns_mcp_error_as_jsonrpc_error():
    from mcp.shared.exceptions import MCPError

    failure = MCPError(code=-32602, message="PRIVATE_MESSAGE", data={"secret": "PRIVATE_DATA"})

    async def call_tool(_name, _arguments):
        raise failure

    registration = ClaudeMcpRegistration(
        {}, {}, session_factory=None, list_tools=None,
    )
    server = registration._server(
        "gateway",
        [("lookup", Tool(name="lookup", input_schema={}), "lookup")],
        SimpleNamespace(call_tool=call_tool),
    )["instance"]
    async with connected_client(server) as client:
        with pytest.raises(MCPError) as raised:
            await client.call_tool("lookup", {})
    assert raised.value.code == failure.code
    assert raised.value.message == failure.message


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", ["ordinary", "result", "cancel", "elicitation", "observer"])
async def test_session_failure_observation_does_not_change_other_outcomes(kind):
    from mcp.shared.exceptions import MCPError, UrlElicitationRequiredError
    from mcp_types import CallToolRequestParams

    observations = []
    callbacks = []
    failure = {
        "ordinary": RuntimeError("timeout permission denied 404"),
        "cancel": asyncio.CancelledError(),
        "elicitation": UrlElicitationRequiredError([]),
        "observer": MCPError(code=-32602, message="PRIVATE"),
    }.get(kind)
    result = CallToolResult(
        content=[TextContent(type="text", text="404 PRIVATE")],
        structured_content={"code": 408},
        is_error=True,
    )

    def observe(item):
        observations.append(item)
        if kind == "observer":
            raise RuntimeError("OBSERVER_PRIVATE")

    def track(callback):
        callbacks.append(callback)
        return callback

    async def call_tool(name, arguments):
        if failure is not None:
            raise failure
        return result

    registration = ClaudeMcpRegistration({}, {}, session_factory=None, list_tools=None, callback_wrapper=track, failure_observer=observe)
    registration._server(
        "gateway",
        [("lookup", Tool(name="lookup", input_schema={}), "lookup")],
        SimpleNamespace(call_tool=call_tool),
    )
    if failure is not None:
        with pytest.raises(type(failure)) as raised:
            await callbacks[0](
                SimpleNamespace(), CallToolRequestParams(name="lookup", arguments={})
            )
        assert raised.value is failure
    else:
        returned = await callbacks[0](
            SimpleNamespace(), CallToolRequestParams(name="lookup", arguments={})
        )
        assert returned.is_error is True
        assert returned.structured_content == result.structured_content
    assert len(observations) == (1 if kind == "observer" else 0)


@pytest.mark.asyncio
@pytest.mark.parametrize("search_mode", [True, "default"])
@pytest.mark.parametrize("integration_kind", ["mcp", "enterprise_gateway"])
async def test_installed_claude_cli_discovers_only_selected_mcp_tools(monkeypatch, tmp_path, search_mode, integration_kind):
    await test_installed_claude_cli_selected_mcp_end_to_end(
        monkeypatch, tmp_path, False, "gateway", integration_kind, tool_search=search_mode,
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("tool_input", [{"query": "lookup"}, {"query": "lookup", "max_results": 1},
    {"query": ""}, {"query": "x", "max_results": True}, {"query": "x", "max_results": 21},
    {"query": "x", "command": "run"}, None])
async def test_sdk_tool_search_is_bounded_and_other_tools_keep_their_admission(tool_input):
    registration = ClaudeMcpRegistration({}, {}, session_factory=None, list_tools=None)
    calls = []
    async def policy(hook_input, *_args):
        calls.append(hook_input)
        return {"ordinary_policy": "denied"}
    matcher = SimpleNamespace(matcher=None, hooks=[policy])
    options = SimpleNamespace(tools=[], allowed_tools=[], hooks={"PreToolUse": [matcher]})
    registration._configure_tool_search(options)
    assert options.tools == options.allowed_tools == ["ToolSearch"]
    callback = matcher.hooks[0]
    result = await callback({"tool_name": "ToolSearch", "tool_input": tool_input})
    valid = tool_input in ({"query": "lookup"}, {"query": "lookup", "max_results": 1})
    assert result["hookSpecificOutput"]["permissionDecision"] == ("allow" if valid else "deny")
    assert calls == []
    assert await callback({"tool_name": "mcp__other__write", "tool_input": {}}) == {"ordinary_policy": "denied"}
    assert len(calls) == 1
